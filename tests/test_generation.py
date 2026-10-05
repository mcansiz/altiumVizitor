"""
@file tests/test_generation.py
@brief Uçtan uca üretim: örnek projelerden TÜM çıktı türleri üretilip denetlenir.

@details Örnekler altium_monkey deposundan sabit commit'le iner
(tools/fetch_samples.py). Her proje için çıktılar `setUpClass`'ta BİR kez
üretilir; bir çıktının hatası diğerlerini gizlemesin diye hatalar saklanıp
ilgili testte raporlanır. Çıktılar `.tmp/test-out/<proje>/` altında kalır
(CI başarısız olursa artifact olarak yüklenir).
"""
import csv
import json
import shutil
import traceback
import unittest

from tests._support import OUT, assert_js_valid, combined_inner, sample_project


class _ProjectBase:
    """Proje başına test kümesi (alt sınıf PROJECT/HAS_PCB/MCU belirler)."""
    PROJECT = None
    HAS_PCB = False
    MCU = None          # MCU pin listesi için designator (None = test atlanır)
    MIN_SHEETS = 1
    MIN_COMPONENTS = 1

    @classmethod
    def setUpClass(cls):
        import viewer
        cls.viewer = viewer
        cls.prj = str(sample_project(cls.PROJECT))
        cls.dir = OUT / cls.PROJECT
        shutil.rmtree(cls.dir, ignore_errors=True)
        cls.dir.mkdir(parents=True)
        cls.logs = []
        cls.errors = {}
        cls.out = {}
        log = cls.logs.append
        jobs = {
            "json": ("generate_json", "proje.json", {}),
            "sch": ("generate_viewer", "sematik.html", {}),
            "pcb": ("generate_pcb_canvas_viewer", "pcb.html", {}),
            "combined": ("generate_combined_viewer", "birlesik.html", {}),
            "bom": ("generate_bom_csv", "bom.csv", {}),
            "pnp": ("generate_pnp_csv", "pnp.csv", {}),
            "icmap": ("generate_ic_map_xlsx", "ic_harita.xlsx", {}),
        }
        if cls.MCU:
            jobs["mcupin"] = ("generate_mcu_pinout_xlsx", "mcu_pin.xlsx",
                              {"mcu_designator": cls.MCU})
        for key, (fn, fname, kw) in jobs.items():
            path = cls.dir / fname
            try:
                getattr(viewer, fn)(cls.prj, str(path), log=log, **kw)
                cls.out[key] = path
            except Exception:
                cls.errors[key] = traceback.format_exc()
        (cls.dir / "uretim.log").write_text("\n".join(map(str, cls.logs)), encoding="utf-8")

    def output(self, key):
        """@brief Üretilmiş çıktının yolunu döndürür; üretim patladıysa testi düşürür."""
        if key in self.errors:
            self.fail(f"{key} üretimi başarısız:\n{self.errors[key]}")
        path = self.out[key]
        self.assertTrue(path.is_file() and path.stat().st_size > 0, f"{path} boş/yok")
        return path

    def html(self, key):
        text = self.output(key).read_text(encoding="utf-8")
        self.assertNotIn("⟪", text, "çeviri işareti kalmış")
        self.assertNotIn("⟫", text, "çeviri işareti kalmış")
        return text

    # --- JSON ---------------------------------------------------------------
    def test_json(self):
        data = json.loads(self.output("json").read_text(encoding="utf-8"))
        s = data["summary"]
        self.assertGreaterEqual(s["sheet_count"], self.MIN_SHEETS)
        self.assertGreaterEqual(s["component_count"], self.MIN_COMPONENTS)
        self.assertTrue(s["has_netlist"], "netlist derlenemedi")
        self.assertGreater(s["pin_connection_count"], 0)
        self.assertEqual(s["has_pnp"], self.HAS_PCB)
        conns = [c for n in data["nets"] for c in n.get("connections", [])]
        self.assertTrue(conns, "hiçbir net'te pin bağlantısı yok")

    # --- HTML görüntüleyiciler ----------------------------------------------
    def test_schematic_viewer(self):
        text = self.html("sch")
        for marker in ("const DRAW_GZ", "SHEET_TREE", 'id="anno-embed"'):
            self.assertIn(marker, text)
        assert_js_valid(self, text, f"{self.PROJECT}_sch")

    def test_schematic_viewer_english(self):
        """İngilizce üretim: dil etiketi `en`, işaret kalıntısı yok, JS geçerli."""
        import i18n
        path = self.dir / "sematik_en.html"
        i18n.set_language("en")
        try:
            self.viewer.generate_viewer(self.prj, str(path), log=lambda _m: None)
        finally:
            i18n.set_language("tr")
        text = path.read_text(encoding="utf-8")
        self.assertIn('lang="en"', text)
        self.assertNotIn("⟪", text)
        self.assertIn("Hierarchy", text)
        assert_js_valid(self, text, f"{self.PROJECT}_sch_en")

    def test_pcb_viewer(self):
        if not self.HAS_PCB:
            self.skipTest("projede PCB yok")
        text = self.html("pcb")
        self.assertIn("const GEO_GZ", text)
        assert_js_valid(self, text, f"{self.PROJECT}_pcb")

    def test_combined_viewer(self):
        text = self.html("combined")
        assert_js_valid(self, text, f"{self.PROJECT}_shell")
        inner = combined_inner(text)
        self.assertIn("SCH", inner)
        if self.HAS_PCB:
            self.assertEqual(set(inner), {"SCH", "PCB", "TD"})
        for name, html in inner.items():
            with self.subTest(panel=name):
                self.assertNotIn("⟪", html)
                if name == "PCB" and not self.HAS_PCB:
                    continue    # "PCB bulunamadı" yer tutucusu: script içermez
                assert_js_valid(self, html, f"{self.PROJECT}_{name}")

    def test_annotations_roundtrip(self):
        """Not taşıma: şematik ve birleşik HTML'e yazılan notlar geri okunuyor."""
        items = [{"id": "t1", "k": "note", "x": 10, "y": 20, "t": "deneme </script> <b>",
                  "fs": 14, "color": "#c62828"},
                 {"id": "t2", "k": "box", "x": 1, "y": 2, "w": 30, "h": 40, "sw": 0.5}]
        for key, where in (("sch", "html"), ("combined", "combined")):
            with self.subTest(target=key):
                src = self.output(key)
                dst = self.dir / f"notlu_{src.name}"
                shutil.copy(src, dst)
                self.assertEqual(self.viewer.write_annotations(dst, items), where)
                self.assertEqual(self.viewer.read_annotations(dst), items)

    # --- Veri çıktıları -----------------------------------------------------
    def test_bom_csv(self):
        with open(self.output("bom"), encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
        self.assertGreater(len(rows), 1, "BOM'da satır yok")

    def test_pnp_csv(self):
        if not self.HAS_PCB:
            self.skipTest("projede PCB yok")
        with open(self.output("pnp"), encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
        self.assertGreater(len(rows), 1, "Pick&Place'te satır yok")

    def test_ic_map_xlsx(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.output("icmap"))
        self.assertGreater(wb.active.max_row, 1)

    def test_mcu_pinout_xlsx(self):
        if not self.MCU:
            self.skipTest("projede MCU yok")
        import openpyxl
        wb = openpyxl.load_workbook(self.output("mcupin"))
        self.assertGreater(wb.active.max_row, 2)


class TestM2Emmc(_ProjectBase, unittest.TestCase):
    """Tek sayfa şematik + PcbDoc: PCB, 3D, BOM/PnP, IC haritası yolu."""
    PROJECT = "m2_emmc"
    HAS_PCB = True
    MCU = "U1"
    MIN_COMPONENTS = 20


class TestSimpleHierarchical(_ProjectBase, unittest.TestCase):
    """PCB'siz hiyerarşik proje: sheet symbol → alt sayfa ağacı yolu."""
    PROJECT = "simple_hierchical"
    HAS_PCB = False
    MIN_SHEETS = 2

    def test_hierarchy_tree(self):
        """Hiyerarşi ağacında bir kök ve en az bir alt düğüm var."""
        import re
        text = self.html("sch")
        m = re.search(r"const SHEET_TREE = (\{.*?\});\n", text, re.S)
        self.assertIsNotNone(m, "SHEET_TREE bulunamadı")
        tree = json.loads(m.group(1))
        nodes = tree["nodes"]
        self.assertTrue(tree["roots"], "kök yok")
        self.assertTrue(any(n["p"] != -1 for n in nodes), "alt sayfa yok")


if __name__ == "__main__":
    unittest.main()
